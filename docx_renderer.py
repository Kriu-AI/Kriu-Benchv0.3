"""
Kriu-Bench-v0.3 DOCX Renderer
===============================
Renders realistic, production-grade .docx resume files across 4 structural layouts:
1. single_col_standard: Classic chronological corporate layout.
2. two_col_sidebar: Modern tech layout with left-hand contact/skills sidebar.
3. tabular_skills_matrix: Structured tables for skills, matrices, and project deliverables.
4. nonstandard_headers: Alternative section naming to stress-test RSSE segmentation.
"""

import os
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex: str):
    """Sets background color of a docx table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text: str, level: int = 1, color_rgb: RGBColor = RGBColor(30, 58, 138)):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = color_rgb
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    return h

def render_docx_single_col_standard(candidate: Dict[str, Any], output_path: str):
    """Layout 1: Classic single-column chronological corporate format."""
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = name_p.add_run(candidate["name"])
    run_name.font.name = "Arial"
    run_name.font.size = Pt(20)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(17, 24, 39)
    name_p.paragraph_format.space_after = Pt(2)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_text = f"{candidate['email']} | {candidate['phone']} | {candidate['location']}"
    if candidate.get("github"):
        contact_text += f" | {candidate['github']}"
    if candidate.get("linkedin"):
        contact_text += f" | {candidate['linkedin']}"
    run_contact = contact_p.add_run(contact_text)
    run_contact.font.name = "Arial"
    run_contact.font.size = Pt(9.5)
    run_contact.font.color.rgb = RGBColor(75, 85, 99)
    contact_p.paragraph_format.space_after = Pt(10)

    # Professional Summary
    add_styled_heading(doc, "PROFESSIONAL SUMMARY")
    sum_p = doc.add_paragraph()
    sum_run = sum_p.add_run(candidate["summary"])
    sum_run.font.name = "Arial"
    sum_run.font.size = Pt(10)
    sum_p.paragraph_format.space_after = Pt(8)

    # Core Skills
    add_styled_heading(doc, "CORE SKILLS & TECHNOLOGIES")
    skills_p = doc.add_paragraph()
    skills_run = skills_p.add_run(" • ".join(candidate["skills_list"]))
    skills_run.font.name = "Arial"
    skills_run.font.size = Pt(9.5)
    skills_p.paragraph_format.space_after = Pt(8)

    # Professional Experience
    add_styled_heading(doc, "PROFESSIONAL EXPERIENCE")
    for job in candidate["experience"]:
        job_header = doc.add_paragraph()
        run_title = job_header.add_run(f"{job['title']} — {job['company']}")
        run_title.font.name = "Arial"
        run_title.font.size = Pt(11)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(31, 41, 55)

        run_dates = job_header.add_run(f" ({job['dates']})")
        run_dates.font.name = "Arial"
        run_dates.font.size = Pt(9.5)
        run_dates.font.italic = True
        run_dates.font.color.rgb = RGBColor(107, 114, 128)
        job_header.paragraph_format.space_after = Pt(2)

        for bullet in job["bullets"]:
            b_p = doc.add_paragraph(style='List Bullet')
            b_run = b_p.add_run(bullet)
            b_run.font.name = "Arial"
            b_run.font.size = Pt(9.5)
            b_p.paragraph_format.space_after = Pt(2)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Projects
    if candidate.get("projects"):
        add_styled_heading(doc, "KEY PROJECTS & TECHNICAL DELIVERABLES")
        for proj in candidate["projects"]:
            p_p = doc.add_paragraph()
            r_proj = p_p.add_run(f"{proj['name']} | Technologies: {proj['tech']}")
            r_proj.font.name = "Arial"
            r_proj.font.size = Pt(10)
            r_proj.font.bold = True
            p_p.paragraph_format.space_after = Pt(1)

            b_p = doc.add_paragraph(style='List Bullet')
            b_run = b_p.add_run(proj["description"])
            b_run.font.name = "Arial"
            b_run.font.size = Pt(9.5)
            b_p.paragraph_format.space_after = Pt(4)

    # Education
    add_styled_heading(doc, "EDUCATION & CREDENTIALS")
    edu_p = doc.add_paragraph()
    edu_run = edu_p.add_run(candidate["education"])
    edu_run.font.name = "Arial"
    edu_run.font.size = Pt(9.5)
    edu_p.paragraph_format.space_after = Pt(4)

    doc.save(output_path)

def render_docx_two_col_sidebar(candidate: Dict[str, Any], output_path: str):
    """Layout 2: Modern tech resume with a 2-column layout (sidebar on left)."""
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Create 1-row, 2-column table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(2.5), Inches(5.0)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    set_cell_background(left_cell, "F3F4F6")
    set_cell_margins(left_cell, top=140, bottom=140, left=180, right=180)
    set_cell_margins(right_cell, top=140, bottom=140, left=200, right=140)

    # Left Sidebar: Name, Contact, Skills, Education
    p_name = left_cell.paragraphs[0]
    r_name = p_name.add_run(candidate["name"])
    r_name.font.name = "Arial"
    r_name.font.size = Pt(16)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(17, 24, 39)
    p_name.paragraph_format.space_after = Pt(6)

    p_contact = left_cell.add_paragraph()
    p_contact.add_run(f"Email:\n{candidate['email']}\n\nPhone:\n{candidate['phone']}\n\nLocation:\n{candidate['location']}\n")
    if candidate.get("github"):
        p_contact.add_run(f"\nGitHub:\n{candidate['github']}\n")
    if candidate.get("linkedin"):
        p_contact.add_run(f"\nLinkedIn:\n{candidate['linkedin']}\n")
    for r in p_contact.runs:
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(55, 65, 81)
    p_contact.paragraph_format.space_after = Pt(12)

    p_sk_h = left_cell.add_paragraph()
    r_sk_h = p_sk_h.add_run("SKILLS")
    r_sk_h.font.name = "Arial"
    r_sk_h.font.size = Pt(11)
    r_sk_h.font.bold = True
    r_sk_h.font.color.rgb = RGBColor(30, 58, 138)
    p_sk_h.paragraph_format.space_after = Pt(4)

    for skill in candidate["skills_list"]:
        sp = left_cell.add_paragraph(style='List Bullet')
        sr = sp.add_run(skill)
        sr.font.name = "Arial"
        sr.font.size = Pt(8.5)
        sp.paragraph_format.space_after = Pt(1)

    p_ed_h = left_cell.add_paragraph()
    r_ed_h = p_ed_h.add_run("\nEDUCATION")
    r_ed_h.font.name = "Arial"
    r_ed_h.font.size = Pt(11)
    r_ed_h.font.bold = True
    r_ed_h.font.color.rgb = RGBColor(30, 58, 138)
    p_ed_h.paragraph_format.space_after = Pt(4)

    p_ed = left_cell.add_paragraph()
    r_ed = p_ed.add_run(candidate["education"])
    r_ed.font.name = "Arial"
    r_ed.font.size = Pt(8.5)

    # Right Column: Summary, Experience, Projects
    p_sum_h = right_cell.paragraphs[0]
    r_sum_h = p_sum_h.add_run("EXECUTIVE SUMMARY")
    r_sum_h.font.name = "Arial"
    r_sum_h.font.size = Pt(12)
    r_sum_h.font.bold = True
    r_sum_h.font.color.rgb = RGBColor(30, 58, 138)
    p_sum_h.paragraph_format.space_after = Pt(3)

    p_sum = right_cell.add_paragraph()
    r_sum = p_sum.add_run(candidate["summary"])
    r_sum.font.name = "Arial"
    r_sum.font.size = Pt(9.5)
    p_sum.paragraph_format.space_after = Pt(10)

    p_exp_h = right_cell.add_paragraph()
    r_exp_h = p_exp_h.add_run("EXPERIENCE")
    r_exp_h.font.name = "Arial"
    r_exp_h.font.size = Pt(12)
    r_exp_h.font.bold = True
    r_exp_h.font.color.rgb = RGBColor(30, 58, 138)
    p_exp_h.paragraph_format.space_after = Pt(4)

    for job in candidate["experience"]:
        p_jh = right_cell.add_paragraph()
        r_jt = p_jh.add_run(f"{job['title']} — {job['company']}")
        r_jt.font.name = "Arial"
        r_jt.font.size = Pt(10)
        r_jt.font.bold = True
        r_jd = p_jh.add_run(f" ({job['dates']})")
        r_jd.font.name = "Arial"
        r_jd.font.size = Pt(9)
        r_jd.font.italic = True
        p_jh.paragraph_format.space_after = Pt(2)

        for b in job["bullets"]:
            bp = right_cell.add_paragraph(style='List Bullet')
            br = bp.add_run(b)
            br.font.name = "Arial"
            br.font.size = Pt(9)
            bp.paragraph_format.space_after = Pt(2)

    if candidate.get("projects"):
        p_pr_h = right_cell.add_paragraph()
        r_pr_h = p_pr_h.add_run("\nPROJECTS")
        r_pr_h.font.name = "Arial"
        r_pr_h.font.size = Pt(12)
        r_pr_h.font.bold = True
        r_pr_h.font.color.rgb = RGBColor(30, 58, 138)
        p_pr_h.paragraph_format.space_after = Pt(4)

        for proj in candidate["projects"]:
            pp = right_cell.add_paragraph()
            rp = pp.add_run(f"{proj['name']} [{proj['tech']}]")
            rp.font.name = "Arial"
            rp.font.size = Pt(9.5)
            rp.font.bold = True
            pp.paragraph_format.space_after = Pt(1)

            pdesc = right_cell.add_paragraph(style='List Bullet')
            rdesc = pdesc.add_run(proj["description"])
            rdesc.font.name = "Arial"
            rdesc.font.size = Pt(9)
            pdesc.paragraph_format.space_after = Pt(3)

    doc.save(output_path)

def render_docx_tabular_skills_matrix(candidate: Dict[str, Any], output_path: str):
    """Layout 3: Structured layout utilizing tables for skills matrix and project delivery."""
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header
    np = doc.add_paragraph()
    nr = np.add_run(candidate["name"])
    nr.font.name = "Arial"
    nr.font.size = Pt(18)
    nr.font.bold = True
    nr.font.color.rgb = RGBColor(17, 24, 39)
    np.paragraph_format.space_after = Pt(2)

    cp = doc.add_paragraph()
    c_text = f"Email: {candidate['email']} | Phone: {candidate['phone']} | Location: {candidate['location']}"
    if candidate.get("github"):
        c_text += f"\nGitHub: {candidate['github']}"
    if candidate.get("linkedin"):
        c_text += f" | LinkedIn: {candidate['linkedin']}"
    cr = cp.add_run(c_text)
    cr.font.name = "Arial"
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(75, 85, 99)
    cp.paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "CAREER SUMMARY")
    sp = doc.add_paragraph()
    sr = sp.add_run(candidate["summary"])
    sr.font.name = "Arial"
    sr.font.size = Pt(9.5)
    sp.paragraph_format.space_after = Pt(8)

    # Technical Skills Matrix Table
    add_styled_heading(doc, "TECHNICAL COMPETENCY MATRIX")
    skills = candidate["skills_list"]
    rows_count = (len(skills) + 2) // 3
    t = doc.add_table(rows=rows_count, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, s in enumerate(skills):
        r_idx = idx // 3
        c_idx = idx % 3
        cell = t.cell(r_idx, c_idx)
        cell.width = Inches(2.3)
        set_cell_background(cell, "F9FAFB")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        run = p.add_run(f"• {s}")
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Experience
    add_styled_heading(doc, "WORK EXPERIENCE")
    for job in candidate["experience"]:
        jp = doc.add_paragraph()
        jr1 = jp.add_run(f"{job['title']} — {job['company']}")
        jr1.font.name = "Arial"
        jr1.font.size = Pt(10.5)
        jr1.font.bold = True
        jr2 = jp.add_run(f" | {job['dates']}")
        jr2.font.name = "Arial"
        jr2.font.size = Pt(9)
        jr2.font.italic = True
        jp.paragraph_format.space_after = Pt(2)

        for b in job["bullets"]:
            bp = doc.add_paragraph(style='List Bullet')
            br = bp.add_run(b)
            br.font.name = "Arial"
            br.font.size = Pt(9)
            bp.paragraph_format.space_after = Pt(2)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Projects Table
    if candidate.get("projects"):
        add_styled_heading(doc, "KEY PROJECT DELIVERABLES")
        pt = doc.add_table(rows=len(candidate["projects"]) + 1, cols=3)
        pt.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Project Name", "Technologies Used", "Impact & Outcomes"]
        for c_idx, h in enumerate(headers):
            cell = pt.cell(0, c_idx)
            set_cell_background(cell, "E5E7EB")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.bold = True

        for r_idx, proj in enumerate(candidate["projects"]):
            row_cells = pt.rows[r_idx + 1].cells
            row_cells[0].width = Inches(1.8)
            row_cells[1].width = Inches(1.8)
            row_cells[2].width = Inches(3.4)

            p0 = row_cells[0].paragraphs[0]
            r0 = p0.add_run(proj["name"])
            r0.font.name = "Arial"; r0.font.size = Pt(8.5); r0.font.bold = True

            p1 = row_cells[1].paragraphs[0]
            r1 = p1.add_run(proj["tech"])
            r1.font.name = "Arial"; r1.font.size = Pt(8.5)

            p2 = row_cells[2].paragraphs[0]
            r2 = p2.add_run(proj["description"])
            r2.font.name = "Arial"; r2.font.size = Pt(8.5)

            for cell in row_cells:
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_styled_heading(doc, "EDUCATION")
    ep = doc.add_paragraph()
    er = ep.add_run(candidate["education"])
    er.font.name = "Arial"
    er.font.size = Pt(9)

    doc.save(output_path)

def render_docx_nonstandard_headers(candidate: Dict[str, Any], output_path: str):
    """Layout 4: Alternative section headers to stress-test RSSE segmentation."""
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header
    np = doc.add_paragraph()
    nr = np.add_run(candidate["name"].upper())
    nr.font.name = "Arial"
    nr.font.size = Pt(17)
    nr.font.bold = True
    nr.font.color.rgb = RGBColor(15, 23, 42)
    np.paragraph_format.space_after = Pt(2)

    cp = doc.add_paragraph()
    c_info = f"{candidate['location']} • {candidate['email']} • {candidate['phone']}"
    if candidate.get("github"):
        c_info += f" • {candidate['github']}"
    if candidate.get("linkedin"):
        c_info += f" • {candidate['linkedin']}"
    cr = cp.add_run(c_info)
    cr.font.name = "Arial"
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(71, 85, 105)
    cp.paragraph_format.space_after = Pt(10)

    # Nonstandard: "CAREER OVERVIEW & PROFILE"
    add_styled_heading(doc, "CAREER OVERVIEW & PROFILE", color_rgb=RGBColor(15, 23, 42))
    sp = doc.add_paragraph()
    sr = sp.add_run(candidate["summary"])
    sr.font.name = "Arial"
    sr.font.size = Pt(9.5)
    sp.paragraph_format.space_after = Pt(8)

    # Nonstandard: "TECHNICAL EXPERTISE & ARSENAL"
    add_styled_heading(doc, "TECHNICAL EXPERTISE & ARSENAL", color_rgb=RGBColor(15, 23, 42))
    skp = doc.add_paragraph()
    skr = skp.add_run(" | ".join(candidate["skills_list"]))
    skr.font.name = "Arial"
    skr.font.size = Pt(9)
    skp.paragraph_format.space_after = Pt(8)

    # Nonstandard: "WORK HISTORY & ACHIEVEMENTS"
    add_styled_heading(doc, "WORK HISTORY & ACHIEVEMENTS", color_rgb=RGBColor(15, 23, 42))
    for job in candidate["experience"]:
        jp = doc.add_paragraph()
        jr = jp.add_run(f"Position: {job['title']} | Organization: {job['company']} ({job['dates']})")
        jr.font.name = "Arial"
        jr.font.size = Pt(10)
        jr.font.bold = True
        jp.paragraph_format.space_after = Pt(2)

        for b in job["bullets"]:
            bp = doc.add_paragraph(style='List Bullet')
            br = bp.add_run(b)
            br.font.name = "Arial"
            br.font.size = Pt(9)
            bp.paragraph_format.space_after = Pt(2)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Nonstandard: "PORTFOLIO & VENTURES"
    if candidate.get("projects"):
        add_styled_heading(doc, "PORTFOLIO & VENTURES", color_rgb=RGBColor(15, 23, 42))
        for proj in candidate["projects"]:
            pp = doc.add_paragraph()
            pr = pp.add_run(f"Initiative: {proj['name']} — Stack: {proj['tech']}")
            pr.font.name = "Arial"
            pr.font.size = Pt(9.5)
            pr.font.bold = True
            pp.paragraph_format.space_after = Pt(1)

            pdesc = doc.add_paragraph(style='List Bullet')
            rdesc = pdesc.add_run(proj["description"])
            rdesc.font.name = "Arial"
            rdesc.font.size = Pt(9)
            pdesc.paragraph_format.space_after = Pt(3)

    # Nonstandard: "ACADEMIC BACKGROUND & CREDENTIALS"
    add_styled_heading(doc, "ACADEMIC BACKGROUND & CREDENTIALS", color_rgb=RGBColor(15, 23, 42))
    ep = doc.add_paragraph()
    er = ep.add_run(candidate["education"])
    er.font.name = "Arial"
    er.font.size = Pt(9)

    doc.save(output_path)

print("DOCX Renderer Loaded Successfully.")
